import hashlib
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from amazon_es_bestseller.collection.candidate_snapshot import parse_saved_candidate_pages
from amazon_es_bestseller.collection.detail_executor import execute_detail_plan
from amazon_es_bestseller.collection.planning import DetailState
from amazon_es_bestseller.monitoring.candidate_detail_handoff import verify_candidate_selection
from amazon_es_bestseller.monitoring.detail_planner import build_detail_plan
from amazon_es_bestseller.quality.source_evidence import audit_detail_run_evidence
from amazon_es_bestseller.cli import main
from amazon_es_bestseller.export.excel import HEAD_ZH


FIXTURES = Path(__file__).parent / "fixtures" / "consolidation_v2"
ASIN = "B013ERLQDK"


def _candidate_snapshot():
    ranking = FIXTURES / "real_ranking_tools.html"
    body = ranking.read_bytes()
    return parse_saved_candidate_pages([{
        "source_type": "BEST_SELLERS",
        "source_url": "https://www.amazon.es/gp/bestsellers/tools/3049288031",
        "page_url": "https://www.amazon.es/gp/bestsellers/tools/3049288031",
        "page_number": 1,
        "html_path": str(ranking),
        "sha256": hashlib.sha256(body).hexdigest(),
        "status_code": 200,
    }], snapshot_id="saved-real-evidence", observed_at="2026-10-01T21:21:23+02:00")


def test_saved_real_candidate_and_detail_evidence_reach_detail_plan_with_traceability(tmp_path):
    snapshot = _candidate_snapshot()
    candidate = next(row for row in snapshot["candidates"]
                     if ASIN in row["candidate_asins_observed"])
    selection = verify_candidate_selection(snapshot, FIXTURES,
                                          candidate_ids=[candidate["candidate_id"]])

    assert snapshot["snapshot_status"] == "UNVERIFIED"
    assert selection["snapshot_status"] == "VERIFIED_SELECTION"
    assert selection["authoritative"] is False
    assert selection["identity_complete"] is False
    plan = build_detail_plan(selection)
    row = next(item for item in plan["records"] if item["asin"] == ASIN)
    assert row["detail_action"] == "REUSE_VALID_CACHE"
    assert row["identity_status"] == "MATCH"
    assert row["candidate_ids"] == [candidate["candidate_id"]]
    assert row["candidate_observation_ids"] == [candidate["observations"][0]["candidate_observation_id"]]

    # Execute the real offline detail runner in pytest's isolated temporary directory.
    output = tmp_path / "detail-run"
    execution = execute_detail_plan(plan, None, str(output), offline=True)
    assert execution["records"][0]["candidate_ids"] == [candidate["candidate_id"]]
    assert execution["records"][0]["candidate_observation_ids"] == [
        candidate["observations"][0]["candidate_observation_id"]]
    audit = audit_detail_run_evidence(output)
    assert audit["evidence_counts"]["VERIFIED_SAVED_EVIDENCE"] == 1
    assert audit["records"][0]["asin"] == ASIN
    assert DetailState(output / "state" / "details_state.json").get(ASIN)["title_es_raw"]


def test_unverified_candidate_and_mismatched_detail_evidence_are_rejected(tmp_path):
    snapshot = _candidate_snapshot()
    candidate = next(row for row in snapshot["candidates"]
                     if ASIN in row["candidate_asins_observed"])
    selection = verify_candidate_selection(snapshot, tmp_path,
                                          candidate_ids=[candidate["candidate_id"]])
    assert selection["verified_candidate_count"] == 0
    with pytest.raises(ValueError, match="未验证"):
        build_detail_plan(selection)

    wrong = tmp_path / f"{ASIN}.html"
    wrong.write_text("<html><body>not a product page</body></html>", encoding="utf-8")
    selection = verify_candidate_selection(snapshot, tmp_path,
                                          candidate_ids=[candidate["candidate_id"]])
    assert selection["verified_candidate_count"] == 0


def test_offline_translation_qa_and_excel_continue_from_verified_detail_run(tmp_path):
    ranking_path = FIXTURES / "real_ranking_tools.html"
    ranking_body = ranking_path.read_bytes()
    candidate_manifest = tmp_path / "candidate-manifest.json"
    candidate_manifest.write_text(json.dumps({
        "snapshot_id": "saved-real-evidence",
        "observed_at": "2026-10-01T21:21:23+02:00",
        "pages": [{
            "source_type": "BEST_SELLERS",
            "source_url": "https://www.amazon.es/gp/bestsellers/tools/3049288031",
            "page_url": "https://www.amazon.es/gp/bestsellers/tools/3049288031",
            "page_number": 1, "html_path": str(ranking_path),
            "sha256": hashlib.sha256(ranking_body).hexdigest(), "status_code": 200,
        }],
    }, ensure_ascii=False), encoding="utf-8")
    candidate_root = tmp_path / "candidate-output"
    assert main(["--offline", "candidate-snapshot", "--manifest", str(candidate_manifest),
                 "--out-dir", str(candidate_root)]) == 0
    snapshot_dir = candidate_root / "candidate_snapshots" / "saved-real-evidence"
    snapshot = {"snapshot_id": "saved-real-evidence",
                "candidates": json.loads((snapshot_dir / "candidates.json").read_text(encoding="utf-8"))}
    candidate = next(row for row in snapshot["candidates"]
                     if ASIN in row["candidate_asins_observed"])
    detail_run = tmp_path / "detail-run"
    selection_dir = tmp_path / "candidate-selection"
    assert main(["candidate-identity-verify", "--snapshot-dir", str(snapshot_dir),
                 "--detail-html-dir", str(FIXTURES), "--candidate-id", candidate["candidate_id"],
                 "--out-dir", str(selection_dir)]) == 0
    selection_manifest = json.loads((selection_dir / "manifest.json").read_text(encoding="utf-8"))
    assert selection_manifest["snapshot_status"] == "VERIFIED_SELECTION"
    assert selection_manifest["authoritative"] is False
    plan_dir = tmp_path / "detail-plan"
    assert main(["detail-plan", "--snapshot", str(selection_dir), "--out-dir", str(plan_dir)]) == 0
    plan_path = plan_dir / "detail_plan.json"
    assert plan_path.is_file()
    assert main(["--offline", "detail-run", "--plan", str(plan_path),
                 "--out-dir", str(detail_run)]) == 0
    execution = json.loads((detail_run / "detail_execution_manifest.json").read_text(encoding="utf-8"))
    assert execution["requested_count"] == 0
    evidence_audit = audit_detail_run_evidence(detail_run)
    assert evidence_audit["evidence_counts"]["VERIFIED_SAVED_EVIDENCE"] == 1

    observation = candidate["observations"][0]
    ranking = {
        "asin": ASIN, "ranking_asin": ASIN,
        "ranking_link_identity_status": "MATCH",
        "ranking_product_url_raw": observation["raw_product_url"],
        "ranking_product_url_normalized": observation["normalized_product_url"],
        "ranking_source_url": observation["source_url"],
        "bestseller_rank": observation["bestseller_rank"],
        "snapshot_id": snapshot["snapshot_id"],
    }
    rankings_path = tmp_path / "rankings.json"
    rankings_path.write_text(json.dumps([ranking], ensure_ascii=False), encoding="utf-8")
    audit_path = tmp_path / "source-audit.json"
    assert main(["source-evidence-audit", "--run-dir", str(detail_run),
                 "--rankings", str(rankings_path), "--out", str(audit_path),
                 "--check", "ranking_identity"]) == 0
    source_audit = json.loads(audit_path.read_text(encoding="utf-8"))
    assert source_audit["quality_gate"]["summary"]["network_requests"] == 0
    assert source_audit["source_audit_compatibility"]["formal_source_audit_pass"] is False

    product = {
        **json.loads((detail_run / "details.json").read_text(encoding="utf-8"))[0],
        "research_category": "DIY与五金",
        "category_l1": "Bricolaje y herramientas",
        "category_l2": "Herramientas manuales y eléctricas",
        "leaf_category": "Herramientas manuales y eléctricas",
        "bestseller_rank": observation["bestseller_rank"],
        "product_url": observation["normalized_product_url"],
    }
    master_path = tmp_path / "master.json"
    master_path.write_text(json.dumps([product], ensure_ascii=False), encoding="utf-8")
    config_path = tmp_path / "translation.json"
    config_path.write_text(json.dumps({"provider_mode": "fake", "fields": ["title_es_raw"]}),
                           encoding="utf-8")
    translation_run = tmp_path / "translation-run"
    common = ["--run-dir", str(translation_run), "--run-id", "offline-closure"]
    assert main(["translation-production", "--stage", "build-input", "--master",
                 str(master_path), *common]) == 0
    assert main(["translation-production", "--stage", "preclean", *common]) == 0
    assert main(["--offline", "translation-production", "--stage", "translate", *common,
                 "--config", str(config_path), "--yes"]) == 0
    qa_path = translation_run / "qa" / "translation_qa.json"
    assert qa_path.is_file()
    assert main(["translation-production", "--stage", "promote", *common,
                 "--config", str(config_path)]) == 0
    workbook_path = tmp_path / "offline-debug.xlsx"
    assert main(["translation-production", "--stage", "export", *common,
                 "--config", str(config_path), "--details", str(detail_run / "details.json"),
                 "--rankings", str(rankings_path), "--html-dir", str(FIXTURES),
                 "--collection-run-dir", str(detail_run), "--debug-export",
                 "--out", str(workbook_path)]) == 0
    workbook = load_workbook(workbook_path, read_only=True)
    assert workbook.sheetnames == ["类目规划", "西班牙语选品清单", "中文选品清单"]
    zh = workbook["中文选品清单"]
    assert [zh.cell(1, column).value for column in range(1, 27)] == HEAD_ZH
    es_asins = {row[1] for row in workbook["西班牙语选品清单"].iter_rows(min_row=2, values_only=True)
                if row[1]}
    zh_asins = {row[2] for row in zh.iter_rows(min_row=2, values_only=True) if row[2]}
    assert es_asins == zh_asins == {ASIN}
    workbook.close()
