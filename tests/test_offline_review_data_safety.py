"""Saved-evidence regressions from the full offline review; no live transport."""
from copy import deepcopy
import hashlib
import json
from pathlib import Path

import openpyxl
import pytest

from amazon_es_bestseller.collection import runtime_io
from amazon_es_bestseller.collection.candidate_snapshot import parse_saved_candidate_pages
from amazon_es_bestseller.collection.detail import CURRENT_DETAIL_SCHEMA_VERSION, collect_details, parse_detail_page
from amazon_es_bestseller.collection.detail_executor import _usable_record, execute_detail_plan
from amazon_es_bestseller.export.excel import export_workbook, merge_manual_fields
from amazon_es_bestseller.monitoring.candidate_detail_handoff import verify_candidate_selection
from amazon_es_bestseller.monitoring.detail_planner import build_detail_plan, validate_detail_plan, write_detail_plan
from amazon_es_bestseller.quality.gate import evaluate_research_readiness

FIXTURES = Path(__file__).parent / 'fixtures'
ASIN = 'B07QHCNQ7M'
OTHER = 'B0B8NPVTC4'


@pytest.fixture(scope='module')
def saved_record():
    html = (FIXTURES / 'html/detail500_20261022' / (ASIN + '.html')).read_text(encoding='utf-8')
    record = parse_detail_page(html, ASIN)
    assert record['title_es_raw']
    return {**record, 'asin': ASIN, 'resolved_asin': ASIN, 'identity_status': 'MATCH',
            'cache_classification': 'VALID_PRODUCT_PAGE', 'access_state': 'NORMAL',
            'detail_schema_version': CURRENT_DETAIL_SCHEMA_VERSION}


def plan(action='REUSE_VALID_CACHE'):
    return {'snapshot_id': 'offline-review', 'records': [
        {'snapshot_id': 'offline-review', 'asin': ASIN, 'detail_action': action}]}


@pytest.mark.parametrize('body', ['{"incomplete":', '{}', '[null]', '[{"asin":"invalid"}]'])
def test_bad_legacy_details_stop_before_overwrite_or_collection(tmp_path, body):
    path = tmp_path / 'details.json'
    path.write_text(body, encoding='utf-8')
    before = path.read_bytes()
    with pytest.raises(ValueError):
        execute_detail_plan(plan('FETCH_NEW'), None, str(tmp_path),
                            collector=lambda *_a, **_k: pytest.fail('must not collect'))
    assert path.read_bytes() == before
    assert not (tmp_path / 'state/details_state.json').exists()
    assert not (tmp_path / 'detail_execution_manifest.json').exists()


def test_legacy_read_permission_error_preserves_file_and_stops(tmp_path, monkeypatch, saved_record):
    path = tmp_path / 'details.json'
    path.write_text(json.dumps([saved_record]), encoding='utf-8')
    before = path.read_bytes()
    original = Path.read_text
    def denied(self, *args, **kwargs):
        if self == path:
            raise PermissionError('injected persistent legacy read denial')
        return original(self, *args, **kwargs)
    monkeypatch.setattr(Path, 'read_text', denied)
    with pytest.raises(PermissionError):
        execute_detail_plan(plan(), None, str(tmp_path), offline=True)
    assert path.read_bytes() == before
    assert not (tmp_path / 'state/details_state.json').exists()


def test_legacy_read_sharing_lock_retries_without_losing_raw_fields(tmp_path, monkeypatch, saved_record):
    path = tmp_path / 'details.json'
    record = {**saved_record, '备注': '保留人工备注', 'unknown_raw_attribute': {'fuente': 'guardada'}}
    path.write_text(json.dumps([record]), encoding='utf-8')
    original = Path.read_text
    calls = []
    def locked_once(self, *args, **kwargs):
        if self == path:
            calls.append(self)
            if len(calls) == 1:
                error = PermissionError('injected sharing lock')
                error.winerror = 32
                raise error
        return original(self, *args, **kwargs)
    monkeypatch.setattr(Path, 'read_text', locked_once)
    monkeypatch.setattr(runtime_io.time, 'sleep', lambda _: None)
    result = execute_detail_plan(plan(), None, str(tmp_path), offline=True)
    assert result['completed_count'] == 1 and result['requested_count'] == 0
    restored = json.loads(path.read_bytes())[0]
    assert restored['title_es_raw'] == record['title_es_raw']
    assert restored['备注'] == record['备注']
    assert restored['unknown_raw_attribute'] == record['unknown_raw_attribute']
    assert len(calls) == 2


@pytest.mark.parametrize('field,value,expected', [
    ('access_state', 'CHALLENGE', 'BLOCK_ACCESS_STATE'),
    ('final_access_state', 'BLOCKED', 'BLOCK_ACCESS_STATE'),
    ('final_access_state', 'RATE_LIMITED', 'BLOCK_ACCESS_STATE'),
    ('detail_status', 'PARSER_ERROR', 'BLOCK_CODE_FIX'),
    ('error_class', 'SCHEMA_ERROR', 'BLOCK_CODE_FIX'),
    ('detail_status', 'NETWORK_ERROR', 'RETRY_TRANSIENT_FAILURE'),
    ('status', 'failed', 'RETRY_TRANSIENT_FAILURE'),
    ('classification', 'INVALID_OR_EMPTY', 'REFETCH_INVALID_CACHE'),
    ('detail_status', 'INVALID', 'REFETCH_INVALID_CACHE'),
    ('is_captcha', True, 'BLOCK_ACCESS_STATE'),
])
def test_failure_evidence_dominates_stale_valid_cache_marker(saved_record, field, value, expected):
    record = {**saved_record, field: value}
    snapshot = {'snapshot_id': 'offline-review', 'snapshot_status': 'AUTHORITATIVE',
                'records': [{'asin': ASIN}]}
    item = build_detail_plan(snapshot, [record])['records'][0]
    assert item['detail_action'] == expected
    assert not _usable_record(record, ASIN)


def test_recovered_initial_challenge_is_not_a_new_access_stop(saved_record):
    record = {**saved_record, 'initial_access_state': 'CHALLENGE', 'final_access_state': 'NORMAL',
              'recovered_from_challenge': True, 'detail_status': 'SUCCESS'}
    snapshot = {'snapshot_id': 'offline-review', 'snapshot_status': 'AUTHORITATIVE',
                'records': [{'asin': ASIN}]}
    assert build_detail_plan(snapshot, [record])['records'][0]['detail_action'] == 'REUSE_VALID_CACHE'
    assert _usable_record(record, ASIN)


def previous_notes(es_rows, zh_rows):
    workbook = openpyxl.Workbook()
    workbook.active.title = '西班牙语选品清单'
    workbook.create_sheet('中文选品清单')
    for name, rows in [('西班牙语选品清单', es_rows), ('中文选品清单', zh_rows)]:
        workbook[name].append(['ASIN', '备注'])
        for row in rows:
            workbook[name].append(row)
    return workbook


def test_notes_are_merged_from_both_sheets_without_modifying_inputs(tmp_path):
    note = '  人工核实\n保留原格式  '
    previous = previous_notes([[ASIN, 'Nota española']], [[OTHER, note]])
    records = [{'asin': ASIN}, {'asin': OTHER}]
    before = deepcopy(records)
    result = merge_manual_fields(records, previous)
    assert {r['asin']: r['备注'] for r in result} == {ASIN: 'Nota española', OTHER: note}
    output = tmp_path / 'offline-notes.xlsx'
    export_workbook(records, prev_workbook=previous, out_path=output)
    workbook = openpyxl.load_workbook(output)
    assert workbook.sheetnames == ['类目规划', '西班牙语选品清单', '中文选品清单']
    assert workbook['中文选品清单'].max_column == 26
    assert {r[2]: r[25] for r in workbook['中文选品清单'].iter_rows(min_row=2, values_only=True)} == {
        ASIN: 'Nota española', OTHER: note}
    assert records == before
    assert previous['中文选品清单'].cell(2, 2).value == note


def test_conflicting_old_notes_stop_export_instead_of_silently_discarding_one(tmp_path):
    previous = previous_notes([[ASIN, '西语表人工备注']], [[ASIN, '中文表另一条人工备注']])
    output = tmp_path / 'must-not-replace.xlsx'
    output.write_bytes(b'prior delivery stays intact')
    with pytest.raises(ValueError, match='MANUAL_NOTES_CONFLICT'):
        export_workbook([{'asin': ASIN}], prev_workbook=previous, out_path=output)
    assert output.read_bytes() == b'prior delivery stays intact'
    assert merge_manual_fields([{'asin': ASIN, '备注': '本次明确保留'}], previous)[0]['备注'] == '本次明确保留'


@pytest.mark.parametrize('second', [ASIN, ASIN.lower()])
def test_duplicate_plan_asins_are_rejected_before_writes(tmp_path, second):
    artifact = plan()
    artifact['records'].append({**artifact['records'][0], 'asin': second})
    with pytest.raises(ValueError, match='DUPLICATE'):
        execute_detail_plan(artifact, None, str(tmp_path), offline=True)
    assert not list(tmp_path.iterdir())


def test_nonrecord_in_hashed_plan_cannot_be_silently_dropped(tmp_path):
    artifact = json.loads(write_detail_plan(plan(), tmp_path)['json'].read_text(encoding='utf-8'))
    artifact['records'].append(None)
    with pytest.raises(ValueError, match='RECORD'):
        validate_detail_plan(artifact)


@pytest.mark.parametrize('bad_asin', ['../escape', '', 'B07QHCNQ7M/extra'])
def test_direct_collector_rejects_bad_identity_before_delivery_or_files(tmp_path, monkeypatch, bad_asin):
    monkeypatch.setattr('amazon_es_bestseller.access.location.ensure_spain_delivery',
                        lambda *_: pytest.fail('invalid identity must stop before delivery navigation'))
    output = tmp_path / 'not-created'
    with pytest.raises(ValueError, match='ASIN'):
        collect_details([bad_asin], None, str(output))
    assert not output.exists()


def test_direct_collector_normalizes_and_deduplicates_saved_asin(tmp_path, monkeypatch):
    root = tmp_path / 'run'
    (root / 'html').mkdir(parents=True)
    (root / 'html' / (ASIN + '.html')).write_bytes(
        (FIXTURES / 'html/detail500_20261022' / (ASIN + '.html')).read_bytes())
    # Same final URL used by the existing real child-page identity regression.
    # This is offline test metadata, never a fabricated network receipt.
    (root / 'html' / (ASIN + '.meta.json')).write_text(json.dumps({
        'status_code': 200, 'final_url': 'https://www.amazon.es/dp/' + ASIN}), encoding='utf-8')
    monkeypatch.setattr('amazon_es_bestseller.access.location.ensure_spain_delivery', lambda *_: None)
    class OfflineSession:
        def goto(self, *_a, **_k):
            pytest.fail('saved duplicate must never be collected again')
        def wait_between_requests(self):
            pass
    progress = []
    rows = collect_details([ASIN, ' ' + ASIN.lower() + ' '], OfflineSession(), str(root),
                           on_progress=progress.append, parser_version='v2')
    assert len(rows) == 1 and rows[0]['asin'] == ASIN
    assert len(progress) == 1 and progress[0]['completed'] == progress[0]['total'] == 1


def test_empty_collector_has_no_delivery_or_file_side_effects(tmp_path, monkeypatch):
    monkeypatch.setattr('amazon_es_bestseller.access.location.ensure_spain_delivery',
                        lambda *_: pytest.fail('empty selection must not navigate'))
    assert collect_details([], None, str(tmp_path)) == []
    assert not list(tmp_path.iterdir())


@pytest.mark.parametrize('records', [None, {'asin': ASIN}, [None]])
def test_invalid_plan_shape_cannot_reduce_the_task_count(records):
    with pytest.raises(ValueError, match='RECORD'):
        validate_detail_plan({'snapshot_id': 'offline-review', 'records': records})


def test_plan_without_header_cannot_mix_snapshot_ids():
    rows = [{**plan()['records'][0], 'snapshot_id': 'a'},
            {**plan()['records'][0], 'asin': OTHER, 'snapshot_id': 'b'}]
    with pytest.raises(ValueError, match='snapshot_id'):
        validate_detail_plan({'records': rows})


def test_explicit_empty_candidate_selection_does_not_select_everything():
    root = FIXTURES / 'consolidation_v2'
    ranking = root / 'real_ranking_tools.html'
    snapshot = parse_saved_candidate_pages([{'source_type': 'BEST_SELLERS',
        'source_url': 'https://www.amazon.es/gp/bestsellers/tools/3049288031',
        'page_url': 'https://www.amazon.es/gp/bestsellers/tools/3049288031',
        'page_number': 1, 'html_path': str(ranking),
        'sha256': hashlib.sha256(ranking.read_bytes()).hexdigest(), 'status_code': 200}],
        snapshot_id='offline-empty-selection', observed_at='2026-10-01T21:21:23+02:00')
    assert verify_candidate_selection(snapshot, root)['verified_candidate_count'] > 0
    result = verify_candidate_selection(snapshot, root, candidate_ids=[])
    assert result['verified_candidate_count'] == 0
    assert result['snapshot_status'] == 'NO_VERIFIED_CANDIDATES'


@pytest.mark.parametrize('checks', [[], [{'check': 'saved-evidence'}],
    [{'check': 'saved-evidence', 'status': 'ERROR'}],
    [{'check': 'saved-evidence', 'status': 'SKIPPED'}],
    [{'status': 'PASS'}]])
def test_quality_gate_cannot_approve_empty_missing_or_unknown_checks(checks):
    result = evaluate_research_readiness(checks)
    assert not result['research_ready']
    assert not result['checks_complete']
    assert result['final_quality_status'] in {'BLOCKED', 'PARTIAL_REVIEW'}
