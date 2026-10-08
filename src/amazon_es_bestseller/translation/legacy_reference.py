"""Conservative, non-Gold reuse candidates from historical Excel exports.

Historical Chinese cells are only references.  They are tied to the current
canonical Spanish field hash and go through the current QA again; this module
never promotes them to a provider-verified translation or a release decision.
"""
from __future__ import annotations

import hashlib
import json
from collections import Counter
from pathlib import Path
from typing import Any, Iterable, Mapping

from .production_contract import (
    PRODUCTION_TRANSLATION_FIELDS,
    canonical_record,
    canonical_source_field,
    source_text,
    target_field_for,
)
from .protection import protect
from .qa import qa_field
from .service import source_hash


LEGACY_REFERENCE_CANDIDATES_SCHEMA_VERSION = "legacy-reference-candidates-v1"
LEGACY_REVIEWED_POLICY_VERSION = "legacy-reviewed-reference-v1"
LEGACY_SOURCE_KIND = "legacy_excel_reference"
LEGACY_PROVIDER = "unknown"

# Only one-cell scalar fields are supported.  Details and bullets require
# item-level boundaries, which a flattened workbook cannot prove.
SUPPORTED_LEGACY_FIELDS = frozenset({
    "title_es_raw", "category_l1", "category_l2", "category_l3",
    "leaf_category", "selected_variation_raw", "specification_es",
    "product_description",
})
UNSUPPORTED_BOUNDARY_FIELDS = frozenset({"product_details", "feature_bullets"})

LEGACY_ES_HEADERS = {
    "商品名称（西语）": "title_es_raw",
    "一级类目": "category_l1",
    "二级类目": "category_l2",
    "三级类目": "category_l3",
    "细分类目": "leaf_category",
    "当前选中规格 / 变体（西语）": "selected_variation_raw",
    "核心规格（西语）": "specification_es",
    "商品描述（西语）": "product_description",
    "产品描述（西语）": "product_description",
}
LEGACY_ZH_HEADERS = {
    "商品名称（中文）": "title_es_raw",
    "一级类目": "category_l1",
    "二级类目": "category_l2",
    "三级类目": "category_l3",
    "细分类目": "leaf_category",
    "当前选中规格 / 变体": "selected_variation_raw",
    "核心规格（中文）": "specification_es",
    "商品描述（中文）": "product_description",
    "产品描述（中文）": "product_description",
}


def _hash_json(value: Any) -> str:
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True,
                                     separators=(",", ":")).encode("utf-8")).hexdigest()


def _canonical_hash(record: Mapping[str, Any]) -> str:
    canonical = canonical_record(record)
    return _hash_json({key: source_text(value) for key, value in canonical.items()})


def _file_hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _read_workbook(path: Path, headers: Mapping[str, str]) -> tuple[dict[str, dict[str, Any]], dict[str, Any]]:
    """Read only the scalar reference cells, retaining exact sheet/row evidence."""
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    if len(workbook.sheetnames) != 1:
        raise ValueError("LEGACY_REFERENCE_WORKBOOK_SHEET_COUNT_UNSAFE")
    sheet = workbook[workbook.sheetnames[0]]
    header_values = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
    indices = {str(value or "").strip(): position for position, value in enumerate(header_values)}
    if "ASIN" not in indices:
        raise ValueError("LEGACY_REFERENCE_ASIN_HEADER_MISSING")
    mapped = {field: indices[name] for name, field in headers.items() if name in indices}
    rows: dict[str, dict[str, Any]] = {}
    for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        asin = str(values[indices["ASIN"]] or "").strip().upper() if indices["ASIN"] < len(values) else ""
        if not asin:
            continue
        if asin in rows:
            raise ValueError("LEGACY_REFERENCE_DUPLICATE_ASIN:%s" % asin)
        rows[asin] = {
            "row": row_number,
            "fields": {field: source_text(values[index]) for field, index in mapped.items()
                       if index < len(values) and source_text(values[index])},
        }
    return rows, {"path": str(path), "file_hash": _file_hash(path), "sheet": sheet.title,
                  "mapped_fields": sorted(mapped.values()), "row_count": len(rows)}


def load_legacy_excel_reference(spanish_path: str | Path, chinese_path: str | Path) -> dict[str, dict[str, Any]]:
    """Pair immutable Spanish/Chinese legacy scalar rows by ASIN.

    This does not assert that the old source matches the current canonical
    record.  The builder performs that required field-hash check later.
    """
    es_rows, es_meta = _read_workbook(Path(spanish_path), LEGACY_ES_HEADERS)
    zh_rows, zh_meta = _read_workbook(Path(chinese_path), LEGACY_ZH_HEADERS)
    references: dict[str, dict[str, Any]] = {}
    for asin in sorted(es_rows.keys() & zh_rows.keys()):
        references[asin] = {
            "spanish": es_rows[asin]["fields"],
            "chinese": zh_rows[asin]["fields"],
            "provenance": {
                "source_kind": LEGACY_SOURCE_KIND,
                "provider": LEGACY_PROVIDER,
                "spanish": {**es_meta, "row": es_rows[asin]["row"]},
                "chinese": {**zh_meta, "row": zh_rows[asin]["row"]},
            },
        }
    return references


def build_legacy_reference_candidates(
        records: Iterable[Mapping[str, Any]], legacy_reference: Mapping[str, Mapping[str, Any]],
        source_gate: Mapping[str, Any]) -> dict[str, Any]:
    """Build field-bound historical candidates without a formal promotion."""
    records = list(records)
    candidates: list[dict[str, Any]] = []
    preserved_source_fields: list[dict[str, Any]] = []
    counts: Counter[str] = Counter()
    for raw_record in records:
        record = canonical_record(raw_record)
        asin = str(record.get("asin") or "").upper()
        canonical_hash = _canonical_hash(raw_record)
        brand = source_text(record.get("brand"))
        if brand:
            preserved_source_fields.append({
                "asin": asin, "field": "brand", "value": brand,
                "reason": "BRAND_PRESERVED_SOURCE_NO_LEGACY_REPLACEMENT",
            })
        legacy = legacy_reference.get(asin)
        if not legacy:
            counts["asin_not_in_legacy_intersection"] += 1
            continue
        legacy_es = legacy.get("spanish") if isinstance(legacy.get("spanish"), Mapping) else {}
        legacy_zh = legacy.get("chinese") if isinstance(legacy.get("chinese"), Mapping) else {}
        for field in PRODUCTION_TRANSLATION_FIELDS:
            field = canonical_source_field(field)
            if field == "brand":
                continue
            if field in UNSUPPORTED_BOUNDARY_FIELDS:
                counts["unsupported_boundary_field"] += 1
                continue
            if field not in SUPPORTED_LEGACY_FIELDS:
                continue
            current_es = source_text(record.get(field))
            historical_es = source_text(legacy_es.get(field))
            historical_zh = source_text(legacy_zh.get(field))
            if not current_es or not historical_es or not historical_zh:
                counts["missing_legacy_field"] += 1
                continue
            field_hash = source_hash(current_es)
            if field_hash != source_hash(historical_es):
                counts["source_hash_mismatch"] += 1
                continue
            qa = qa_field(protect(current_es, protected_values=(brand,)), historical_zh, current_es,
                          field=field, brand=brand)
            candidate_status = "PASS" if qa["qa_status"] == "pass" else "QA_FAILED"
            candidates.append({
                "asin": asin,
                "field": field,
                "target_field": target_field_for(field),
                "legacy_es": historical_es,
                "legacy_zh": historical_zh,
                "field_hash": field_hash,
                "current_canonical_hash": canonical_hash,
                "source_kind": LEGACY_SOURCE_KIND,
                "provider": LEGACY_PROVIDER,
                "provenance": legacy.get("provenance") or {},
                "qa": qa,
                "candidate_status": candidate_status,
                "release_admission": "DIAGNOSTIC_ONLY_SOURCE_GATE_BLOCKED"
                if str(source_gate.get("status") or "").upper() != "SOURCE_READY"
                else "DIAGNOSTIC_ONLY_LEGACY_REVIEW_REQUIRED",
            })
            counts[candidate_status.lower()] += 1
    current_canonical_hash = _hash_json(sorted(_canonical_hash(record) for record in records))
    qa_issue_rows = [
        {"asin": candidate["asin"], "field": candidate["field"], **issue}
        for candidate in candidates for issue in candidate["qa"].get("issues", [])
    ]
    return {
        "schema_version": LEGACY_REFERENCE_CANDIDATES_SCHEMA_VERSION,
        "status": "DIAGNOSTIC_ONLY",
        "source_gate_status": str(source_gate.get("status") or "UNKNOWN"),
        "source_binding": {"current_canonical_hash": current_canonical_hash},
        "candidates": candidates,
        "qa_report": {
            "schema_version": "legacy-reference-qa-report-v1",
            "status": "pass" if not qa_issue_rows else "qa_failed",
            "counts": {
                "candidates": len(candidates),
                "pass": counts["pass"],
                "qa_failed": counts["qa_failed"],
                "issues": len(qa_issue_rows),
            },
            "issues": qa_issue_rows,
            "diagnostic_only": True,
        },
        "preserved_source_fields": preserved_source_fields,
        "summary": dict(sorted(counts.items())),
        "formal_promotion": {
            "policy_version_required": LEGACY_REVIEWED_POLICY_VERSION,
            "requires_source_ready_reverification": True,
            "requires_explicit_candidate_review": True,
            "legacy_is_not_provider_verified": True,
        },
    }


def prepare_legacy_review_input(
        candidate_manifest: Mapping[str, Any], *, source_candidate_manifest_path: str | Path,
        source_candidate_manifest_hash: str, available_asins: Iterable[object],
        selected_asins: Iterable[str] | None = None) -> dict[str, Any]:
    """Bind pending semantic review to real producer files, never a synthetic gate.

    Automatic QA is rechecked but cannot create KEEP or write translation memory.
    The existing source loader also verifies independent raw-parent authority.
    """
    from ..orchestration.translation_batch import load_source_candidate

    path = Path(source_candidate_manifest_path)
    loaded = load_source_candidate(path, manifest_hash=source_candidate_manifest_hash,
                                   available_asins=available_asins)
    manifest_bytes = path.read_bytes()
    manifest = json.loads(manifest_bytes)
    master_path = path.parent / "spanish_master_5480.json"
    master_bytes = master_path.read_bytes()
    if (hashlib.sha256(manifest_bytes).hexdigest() != source_candidate_manifest_hash
            or hashlib.sha256(master_bytes).hexdigest() != manifest["artifacts"].get(master_path.name)):
        raise ValueError("LEGACY_SOURCE_CHANGED_AFTER_VERIFICATION")
    producer = json.loads(master_bytes)
    records = producer["records"]
    current_hash = _hash_json(sorted(_canonical_hash(row) for row in records))
    if (candidate_manifest.get("schema_version") != LEGACY_REFERENCE_CANDIDATES_SCHEMA_VERSION
            or candidate_manifest.get("source_binding", {}).get("current_canonical_hash") != current_hash):
        raise ValueError("LEGACY_CURRENT_CONTEXT_HASH_MISMATCH")
    by_asin = {row["asin"]: row for row in records}
    selected = set(selected_asins) if selected_asins is not None else set(by_asin)
    if not selected or not selected <= set(by_asin):
        raise ValueError("LEGACY_REVIEW_ASIN_SCOPE_MISMATCH")
    output, seen = [], set()
    for candidate in candidate_manifest.get("candidates") or []:
        asin, field = str(candidate.get("asin") or ""), str(candidate.get("field") or "")
        if (asin not in by_asin or field not in SUPPORTED_LEGACY_FIELDS or (asin, field) in seen
                or candidate.get("provider") != LEGACY_PROVIDER or candidate.get("source_kind") != LEGACY_SOURCE_KIND
                or candidate.get("target_field") != target_field_for(field)):
            raise ValueError("LEGACY_CANDIDATE_IDENTITY_OR_PROVENANCE_INVALID")
        seen.add((asin, field))
        if asin not in selected:
            continue
        record = canonical_record(by_asin[asin])
        source, translated = source_text(record.get(field)), source_text(candidate.get("legacy_zh"))
        context_hash = _canonical_hash(by_asin[asin])
        if (not source or not translated or source != candidate.get("legacy_es")
                or source_hash(source) != candidate.get("field_hash")
                or context_hash != candidate.get("current_canonical_hash")):
            raise ValueError("LEGACY_CANDIDATE_CURRENT_FIELD_BINDING_INVALID")
        brand = source_text(record.get("brand"))
        qa = qa_field(protect(source, protected_values=(brand,)), translated, source, field=field, brand=brand)
        output.append({"asin": asin, "field": field, "target_field": target_field_for(field),
            "source_value": source, "source_hash": source_hash(source), "candidate_value": translated,
            "candidate_hash": _hash_json(candidate), "context_hash": context_hash, "context": record,
            "qa": qa, "decision": "PENDING_SEMANTIC_REVIEW", "semantic_review_required": True,
            "provider": LEGACY_PROVIDER, "source_kind": LEGACY_SOURCE_KIND,
            "resolution_source": "legacy-reviewed-reference", "provenance": candidate.get("provenance") or {}})
    authority = {"source_candidate_manifest_path": str(path.resolve()),
        "source_candidate_manifest_hash": loaded["source_candidate_manifest_hash"],
        "source_master_file_hash": hashlib.sha256(master_bytes).hexdigest(),
        "current_canonical_hash": current_hash, "source_gate_audit_hash": producer["source_gate"]["audit_hash"],
        "selected_asins": sorted(selected)}
    return {"schema_version": "legacy-verified-review-input-v1", "status": "LEGACY_PENDING_SEMANTIC_REVIEW",
        "authority": authority, "candidates": output, "admitted_candidate_hashes": [],
        "provider_calls": 0, "formal_tm_writes": 0, "automatic_promotions": 0}


def evaluate_legacy_reviewed_formal_gate(
        candidate_manifest: Mapping[str, Any], source_gate: Mapping[str, Any],
        policy_evidence: Mapping[str, Any], *, source_candidate_manifest_path: str | Path | None = None,
        source_candidate_manifest_hash: str = "", available_asins: Iterable[object] = (),
        selected_asins: Iterable[str] | None = None) -> dict[str, Any]:
    """Require independently loaded current authority and explicit semantic KEEP."""
    if str(source_gate.get("status") or "").upper() != "SOURCE_READY":
        return {"status": "SOURCE_GATE_NOT_READY", "admitted_candidate_hashes": []}
    if source_candidate_manifest_path is None:
        return {"status": "SOURCE_BINDING_REVERIFY_REQUIRED", "admitted_candidate_hashes": []}
    review = prepare_legacy_review_input(candidate_manifest,
        source_candidate_manifest_path=source_candidate_manifest_path,
        source_candidate_manifest_hash=source_candidate_manifest_hash,
        available_asins=available_asins, selected_asins=selected_asins)
    if source_gate.get("audit_hash") != review["authority"]["source_gate_audit_hash"] or not source_gate.get("ready"):
        return {"status": "SOURCE_BINDING_REVERIFY_REQUIRED", "admitted_candidate_hashes": []}
    if policy_evidence.get("policy_version") != LEGACY_REVIEWED_POLICY_VERSION:
        return {"status": "LEGACY_REVIEW_POLICY_MISSING", "admitted_candidate_hashes": []}
    decisions = policy_evidence.get("semantic_decisions") or []
    reviewed = {str(item.get("candidate_hash") or ""): item for item in decisions if isinstance(item, Mapping)}
    if len(reviewed) != len(decisions):
        raise ValueError("LEGACY_SEMANTIC_DECISION_IDENTITY_INVALID")
    admitted = []
    for candidate in review["candidates"]:
        decision = reviewed.get(candidate["candidate_hash"]) or {}
        if (candidate["qa"]["qa_status"] == "pass" and decision.get("decision") == "KEEP"
                and decision.get("semantic_review_status") == "PASS"
                and decision.get("review_model") in {"CODEX", "OWNER"} and str(decision.get("review_note") or "").strip()
                and decision.get("source_hash") == candidate["source_hash"]
                and decision.get("context_hash") == candidate["context_hash"]
                and decision.get("reviewed_value") == candidate["candidate_value"]):
            admitted.append(candidate["candidate_hash"])
    return {"status": "LEGACY_REVIEWED_REFERENCE_READY" if admitted and len(admitted) == len(review["candidates"])
            else "LEGACY_REVIEW_REQUIRED", "admitted_candidate_hashes": admitted,
            "pending_review_count": len(review["candidates"]) - len(admitted), "formal_tm_writes": 0}


def admit_legacy_run_reference_cache(
        candidate_manifest: Mapping[str, Any], source_gate: Mapping[str, Any],
        policy_evidence: Mapping[str, Any], *, output_path: str | Path,
        dry_run: bool = False, **authority: Any) -> dict[str, Any]:
    """Write only exact, formally admitted fields to a new ASIN-scoped run cache.

    This is an agent/reference namespace, never Qwen evidence, human gold or
    global TM. A partially reviewed batch retains its overall review gate.
    """
    from .cache import TranslationCache

    path = Path(output_path)
    if path.exists():
        raise FileExistsError("LEGACY_RUN_CACHE_ALREADY_EXISTS")
    gate = evaluate_legacy_reviewed_formal_gate(candidate_manifest, source_gate, policy_evidence, **authority)
    if gate["status"] not in {"LEGACY_REVIEW_REQUIRED", "LEGACY_REVIEWED_REFERENCE_READY"}:
        raise ValueError("LEGACY_FORMAL_GATE_BLOCKED:" + gate["status"])
    requested = {str(row.get("candidate_hash") or "") for row in policy_evidence.get("semantic_decisions") or []}
    admitted = set(gate["admitted_candidate_hashes"])
    if requested != admitted:
        raise ValueError("LEGACY_REQUESTED_ADMISSION_NOT_VERIFIED")
    review = prepare_legacy_review_input(candidate_manifest, **authority)
    decisions = {row["candidate_hash"]: row for row in policy_evidence.get("semantic_decisions") or []}
    envelopes = []
    for row in review["candidates"]:
        if row["candidate_hash"] not in admitted:
            continue
        decision = decisions[row["candidate_hash"]]
        envelopes.append({"asin": row["asin"], "field": row["field"], "target_field": row["target_field"],
            "source_text": row["source_value"], "source_hash": row["source_hash"],
            "translated_text": row["candidate_value"], "candidate_text": row["candidate_value"],
            "translation_status": "success", "qa_status": "pass", "qa_issues": [],
            "provider": LEGACY_PROVIDER, "model": "legacy-reviewed-reference",
            "source_kind": LEGACY_SOURCE_KIND, "resolution_source": "legacy-reviewed-reference",
            "schema_version": LEGACY_REVIEWED_POLICY_VERSION, "prompt_version": "explicit-review-v1",
            "attempt_count": 0, "last_error": None, "context_hash": row["context_hash"],
            "candidate_hash": row["candidate_hash"], "review_model": decision["review_model"],
            "review_origin": "agent" if decision["review_model"] == "CODEX" else "human",
            "review_note": decision["review_note"], "semantic_review_status": "PASS",
            "human_gold": False, "provider_verified": False, "formal_authority": review["authority"],
            "review_artifact_sha256": policy_evidence.get("review_artifact_sha256", ""),
            "provenance": row["provenance"]})
    report = {"status": "DRY_ADMISSION_VERIFIED" if dry_run else "RUN_REFERENCES_ADMITTED",
        "gate": gate, "admitted_fields": len(envelopes), "held_fields": len(review["candidates"]) - len(envelopes),
        "admitted_candidate_hashes": sorted(admitted), "cache_path": str(path),
        "cache_written": False, "provider_calls": 0, "formal_tm_writes": 0, "automatic_promotions": 0,
        "authority": review["authority"], "reference_envelopes": envelopes}
    if not dry_run and envelopes:
        cache = TranslationCache(path)
        for envelope in envelopes:
            key = cache.key(envelope["asin"], envelope["field"], envelope["source_hash"], LEGACY_PROVIDER,
                            "legacy-reviewed-reference", LEGACY_REVIEWED_POLICY_VERSION, "explicit-review-v1")
            cache.put(key, envelope)
        cache.save()
        report["cache_written"] = True
    return report
