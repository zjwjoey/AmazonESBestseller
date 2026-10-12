"""Read-only XLSX/data validation; write a new verification receipt only."""
from __future__ import annotations

import argparse
import hashlib
import json
import math
from collections import Counter
from pathlib import Path
from zipfile import ZipFile
from itertools import zip_longest

import openpyxl


def read(path):
    return json.loads(path.read_text(encoding="utf-8"))


def sha(path):
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", required=True)
    args = parser.parse_args()
    output = Path(args.output_dir).resolve()
    review = output / "final_review"
    payload = read(review / "workbook_payload.json")
    book = review / "AmazonES_5480SKU_西语离线回归候选_20261009.xlsx"
    wb = openpyxl.load_workbook(book, read_only=True, data_only=False)
    assert wb.sheetnames == [s["name"] for s in payload["sheets"]]
    cells_checked, escaped_strings, mismatches, errors = 0, 0, [], []
    sheet_counts = {}
    error_values = {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!", "#SPILL!", "#CALC!"}
    for spec in payload["sheets"]:
        sheet = wb[spec["name"]]
        expected_rows = [spec["headers"], *spec["rows"]]
        # XLSX worksheet dimensions are optional; validate actual rows/cells, not a metadata hint.
        for row_index, (actual, expected) in enumerate(zip_longest(
            sheet.iter_rows(max_col=len(spec["headers"])), expected_rows), 1):
            assert actual is not None and expected is not None, (spec["name"], row_index)
            for col_index, (cell, target) in enumerate(zip(actual, expected), 1):
                cells_checked += 1
                value = cell.value
                if spec["name"] == "类目规划" and row_index == 2 and col_index == 2:
                    assert value == "=COUNTA('西班牙语选品清单'!B2:B5481)"
                    continue
                if isinstance(target, str) and target.startswith(("=", "+", "@")):
                    escaped_strings += 1
                    target = "'" + target
                if target == "":
                    target = None
                same = value == target
                if not same and isinstance(value, (int, float)) and isinstance(target, (int, float)):
                    same = math.isclose(value, target, rel_tol=1e-12, abs_tol=1e-12)
                if not same:
                    mismatches.append({"sheet": spec["name"], "row": row_index, "column": col_index})
                if cell.data_type == "e" or (cell.data_type == "f" and value in error_values):
                    errors.append({"sheet": spec["name"], "coordinate": cell.coordinate, "value": value})
        sheet_counts[spec["name"]] = len(spec["rows"])
    wb.close()
    cached = openpyxl.load_workbook(book, read_only=True, data_only=True)
    assert cached["类目规划"]["B2"].value == 5480
    es = [row[1] for row in cached["西班牙语选品清单"].iter_rows(min_row=2, values_only=True)]
    zh = [row[2] for row in cached["中文选品清单"].iter_rows(min_row=2, values_only=True)]
    records = read(output / "spanish_candidate/spanish_master_5480.json")["records"]
    asins = {r["asin"] for r in records}
    assert len(es) == len(zh) == len(asins) == 5480
    assert len(set(es)) == len(set(zh)) == 5480 and set(es) == set(zh) == asins
    assert all(isinstance(v, str) for v in es + zh)
    cached.close()
    metadata = read(output / "input_metadata.json")
    hashes = {name: sha(Path(path)) == metadata["input_sha256"][name]
              for name, path in metadata["inputs"].items()}
    code_hashes = {name: sha(output.parents[1] / name) == expected
                   for name, expected in metadata["code_file_sha256"].items()}
    assert all(hashes.values()) and all(code_hashes.values())
    summary = read(review / "summary.json")
    sku_results = read(review / "sku_processing_results.json")
    integrity_issues = [r for r in read(review / "unresolved.json") if r.get("origin") == "final_integrity"]
    conflicts = {r["asin"] for r in sku_results if r["detail_attribute_asin_conflicts"]}
    assert {r["asin"] for r in integrity_issues} == conflicts
    with ZipFile(book) as package:
        native_tables = [name for name in package.namelist() if name.startswith("xl/tables/table") and name.endswith(".xml")]
        native_sheets = [name for name in package.namelist() if name.startswith("xl/worksheets/sheet") and name.endswith(".xml")]
        assert len(native_tables) == 6 and len(native_sheets) == 7
        assert all(b'pane' in package.read(name) and b'ySplit="1"' in package.read(name) for name in native_sheets)
    assert not mismatches and not errors, {"mismatches": mismatches[:10], "errors": errors[:10]}
    receipt = {"delivery_validation_passed": True, "workbook_path": str(book), "workbook_sha256": sha(book),
        "cells_compared_exactly_or_numeric_tolerance": cells_checked, "formula_cached_count": 5480,
        "escaped_formula_like_text": escaped_strings, "mismatch_count": len(mismatches), "formula_error_count": len(errors),
        "sheet_rows": sheet_counts, "unique_asins": 5480, "all_asins_text": True,
        "input_hashes_unchanged": hashes, "source_code_hashes_unchanged": all(code_hashes.values()),
        "identity_conflict_skus": len(conflicts), "unique_skus_current_or_identity_issue": len(conflicts | {
            i["asin"] for i in read(output / "spanish_candidate/audit.json")["closure_audit"]["issues"] if i.get("asin")}),
        "historical_comparison": summary["comparison_by_layer"], "actual_raw_field_changes": summary["actual_raw_field_changes"],
        "current_issue_types": summary["current_issue_types"], "source_gate_status": summary["source_gate"]["status"],
        "formal_use_allowed": False, "visual_review": "Preview images retained; reviewed by agent separately",
        "field_empty_counts": dict(Counter(key for row in records for key in (
            "title_es_raw", "brand", "category_l1", "category_l2", "category_l3", "leaf_category", "product_details_es",
            "feature_bullets_raw", "specification_es") if not row.get(key)))}
    with (review / "delivery_verification.json").open("x", encoding="utf-8") as stream:
        json.dump(receipt, stream, ensure_ascii=False, indent=2)
    print(json.dumps(receipt, ensure_ascii=False))


if __name__ == "__main__":
    main()
