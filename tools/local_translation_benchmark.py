"""Read-only local ES/ZH translation benchmark for existing research files.

This is intentionally outside production runtime.  It compares existing local
Spanish and Chinese exports by ASIN and by one canonical field at a time; it
does not call a provider, rewrite a workbook, or promote dictionary entries.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any, Iterable, Mapping

from amazon_es_bestseller.quality.chinese import audit_field
from amazon_es_bestseller.translation.production_contract import (
    PRODUCTION_FIELD_ALIASES, canonical_source_field, source_text, target_field_for,
)
from amazon_es_bestseller.translation.service import source_hash


TARGET_ALIASES = {
    "title_zh": ("title_zh", "\u5546\u54c1\u540d\u79f0\uff08\u4e2d\u6587\uff09"),
    "category_l1_zh": ("category_l1_zh", "\u4e00\u7ea7\u7c7b\u76ee\uff08\u4e2d\u6587\uff09", "\u4e00\u7ea7\u7c7b\u76ee"),
    "category_l2_zh": ("category_l2_zh", "\u4e8c\u7ea7\u7c7b\u76ee\uff08\u4e2d\u6587\uff09", "\u4e8c\u7ea7\u7c7b\u76ee"),
    "category_l3_zh": ("category_l3_zh", "\u4e09\u7ea7\u7c7b\u76ee\uff08\u4e2d\u6587\uff09", "\u4e09\u7ea7\u7c7b\u76ee"),
    "leaf_category_zh": ("leaf_category_zh", "\u7ec6\u5206\u7c7b\u76ee\uff08\u4e2d\u6587\uff09", "\u7ec6\u5206\u7c7b\u76ee"),
    "selected_variation_zh": ("selected_variation_zh", "\u5f53\u524d\u9009\u4e2d\u89c4\u683c / \u53d8\u4f53\uff08\u4e2d\u6587\uff09", "\u5f53\u524d\u9009\u4e2d\u89c4\u683c / \u53d8\u4f53"),
    "specification_zh": ("specification_zh", "\u6838\u5fc3\u89c4\u683c\uff08\u4e2d\u6587\uff09"),
    "product_details_zh": ("product_details_zh", "\u5b8c\u6574\u5546\u54c1\u8be6\u60c5\uff08\u4e2d\u6587\uff09"),
    "feature_bullets_zh": ("feature_bullets_zh", "\u5546\u54c1\u5356\u70b9\uff08\u4e2d\u6587\uff09"),
    "description_zh": ("description_zh", "\u5546\u54c1\u63cf\u8ff0\uff08\u4e2d\u6587\uff09"),
}
ASIN_ALIASES = ("asin", "ASIN", "\u5546\u54c1 ASIN")
P0_CODES = {"SOURCE_HASH_MISMATCH", "PROTECTED_TOKEN_MISSING", "NEGATION_MISMATCH", "PRODUCT_TYPE_ERROR"}
P1_CODES = {"NUMERIC_MISMATCH", "UNIT_MISMATCH", "EMPTY_TRANSLATION", "UNSUPPORTED_CLAIM"}
P2_CODES = {"SPANISH_RESIDUAL", "UNREADABLE_OR_MARKUP", "SOURCE_HASH_UNVERIFIABLE",
            "SOURCE_EMPTY_TARGET_PRESENT", "TARGET_FIELD_MISSING"}
TOKEN_RE = re.compile(r"[A-Za-z][A-Za-z0-9+.-]{2,}")


def _hash_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _header_key(value: Any) -> str:
    return re.sub(r"\s+", "", str(value or "")).casefold()


def _row_value(row: Mapping[str, Any], aliases: Iterable[str]) -> str:
    indexed = {_header_key(key): value for key, value in row.items()}
    for alias in aliases:
        value = indexed.get(_header_key(alias))
        if value not in (None, ""):
            return str(value).strip()
    return ""


def _read_rows(path: Path) -> tuple[list[dict[str, Any]], list[str]]:
    if path.suffix.casefold() == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle)
            return [dict(row) for row in reader], list(reader.fieldnames or [])
    if path.suffix.casefold() == ".xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.worksheets[0]
        iterator = sheet.iter_rows(values_only=True)
        headers = [str(value or "") for value in next(iterator, ())]
        rows = [{headers[index]: values[index] if index < len(values) else None
                 for index in range(len(headers))} for values in iterator]
        return rows, headers
    raise ValueError("UNSUPPORTED_INPUT:%s" % path.suffix)


def _index_rows(rows: Iterable[Mapping[str, Any]], source: str) -> tuple[dict[str, dict[str, Any]], list[dict[str, Any]]]:
    index: dict[str, dict[str, Any]] = {}
    issues: list[dict[str, Any]] = []
    for offset, row in enumerate(rows, 2):
        asin = _row_value(row, ASIN_ALIASES).upper()
        if not asin:
            issues.append({"severity": "P0", "code": "ASIN_MISSING", "dataset": source, "row": offset})
        elif asin in index:
            issues.append({"severity": "P0", "code": "ASIN_DUPLICATE", "dataset": source, "asin": asin, "row": offset})
        else:
            index[asin] = dict(row)
    return index, issues


def _source_aliases(field: str) -> tuple[str, ...]:
    return tuple(PRODUCTION_FIELD_ALIASES.get(field, (field,)))


def _severity(code: str) -> str:
    if code in P0_CODES:
        return "P0"
    if code in P1_CODES:
        return "P1"
    return "P2"


def _issue(asin: str, field: str, code: str, *, source: str = "", target: str = "", detail: Any = None) -> dict[str, Any]:
    return {"asin": asin, "field": field, "target_field": target_field_for(field), "severity": _severity(code),
            "code": code, "source_text": source, "translated_text": target, "detail": detail or {}}


def _write_csv(path: Path, rows: list[Mapping[str, Any]]) -> None:
    headers = sorted({key for row in rows for key in row} or {"status"})
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers, extrasaction="ignore")
        writer.writeheader(); writer.writerows(rows)


def benchmark(es_path: Path, zh_path: Path, output: Path) -> dict[str, Any]:
    output.mkdir(parents=True, exist_ok=True)
    before = {"es": _hash_file(es_path), "zh": _hash_file(zh_path)}
    es_rows, es_headers = _read_rows(es_path)
    zh_rows, zh_headers = _read_rows(zh_path)
    es_by_asin, issues = _index_rows(es_rows, "es")
    zh_by_asin, index_issues = _index_rows(zh_rows, "zh")
    issues.extend(index_issues)
    source_fields = [field for field in PRODUCTION_FIELD_ALIASES if target_field_for(field) in TARGET_ALIASES]
    common = sorted(set(es_by_asin) & set(zh_by_asin))
    for asin in sorted(set(es_by_asin) - set(zh_by_asin)):
        issues.append(_issue(asin, "", "TARGET_RECORD_MISSING"))
    for asin in sorted(set(zh_by_asin) - set(es_by_asin)):
        issues.append(_issue(asin, "", "SOURCE_RECORD_MISSING"))
    outcomes: list[dict[str, Any]] = []
    dictionary: dict[tuple[str, str, str], set[str]] = defaultdict(set)
    for asin in common:
        source_row, target_row = es_by_asin[asin], zh_by_asin[asin]
        context = _row_value(source_row, _source_aliases("leaf_category")) or _row_value(source_row, _source_aliases("category_l3"))
        for field in source_fields:
            target_field = target_field_for(field)
            source = _row_value(source_row, _source_aliases(field))
            target = _row_value(target_row, TARGET_ALIASES[target_field])
            if not source and not target:
                continue
            if not source and target:
                item = _issue(asin, field, "SOURCE_EMPTY_TARGET_PRESENT", target=target)
                issues.append(item); outcomes.append({**item, "status": "WARN"}); continue
            if source and not target:
                item = _issue(asin, field, "EMPTY_TRANSLATION", source=source)
                issues.append(item); outcomes.append({**item, "status": "REPAIR"}); continue
            # Existing display exports normally have no field-level source hash.
            # Analyze content but never elevate the row to binding-verified PASS.
            qa = audit_field(asin=asin, field=field, target_field=target_field, field_type=target_field,
                             source_es=source, translated_zh=target, source_hash=source_hash(source),
                             context={"category": context, "field": field, "target_field": target_field})
            item = {"asin": asin, "field": field, "target_field": target_field, "source_text": source,
                    "translated_text": target, "status": qa["status"], "issues": qa["issues"],
                    "binding": "SOURCE_HASH_UNVERIFIABLE"}
            outcomes.append(item)
            issues.append(_issue(asin, field, "SOURCE_HASH_UNVERIFIABLE", source=source, target=target))
            for issue in qa["issues"]:
                issues.append(_issue(asin, field, str(issue.get("code") or "QA_ISSUE"), source=source,
                                     target=target, detail=issue))
            if qa["status"] == "PASS":
                normalized = source.casefold().strip()
                dictionary[(field, context.casefold(), normalized)].add(target)
    after = {"es": _hash_file(es_path), "zh": _hash_file(zh_path)}
    if before != after:
        raise RuntimeError("INPUT_HASH_CHANGED_DURING_READ")
    field_counts: dict[str, Counter] = defaultdict(Counter)
    for row in outcomes:
        field_counts[row.get("field") or "record"][row.get("status") or "WARN"] += 1
    issue_counts = Counter((row["severity"], row["code"]) for row in issues)
    suspicious = [row for row in outcomes if row.get("status") != "PASS"]
    high_errors = [row for row in issues if row["severity"] in {"P0", "P1"}]
    high_pass = [row for row in outcomes if row.get("status") == "PASS" and not row.get("issues")]
    candidates = [{"field": field, "context": context, "source_term": source, "target_variants": sorted(values),
                   "frequency": sum(1 for row in outcomes if row.get("field") == field and row.get("source_text", "").casefold().strip() == source),
                   "promotion_status": "CANDIDATE_ONLY", "reason": "LOCAL_ZH_NOT_GOLD"}
                  for (field, context, source), values in dictionary.items()]
    conflicts = [row for row in candidates if len(row["target_variants"]) > 1]
    summary = {"status": "COMPLETE", "authority": "LOCAL_ZH_NOT_GOLD", "input": {"es": str(es_path), "zh": str(zh_path),
               "hash_before": before, "hash_after": after, "es_headers": es_headers, "zh_headers": zh_headers},
               "unique_skus": {"es": len(es_by_asin), "zh": len(zh_by_asin), "matched": len(common)},
               "recognized_field_pairs": len(source_fields), "outcomes": dict(Counter(row["status"] for row in outcomes)),
               "issues": {"P0": sum(1 for row in issues if row["severity"] == "P0"),
                          "P1": sum(1 for row in issues if row["severity"] == "P1"),
                          "P2": sum(1 for row in issues if row["severity"] == "P2")}}
    (output / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    _write_csv(output / "field_summary.csv", [{"field": field, **counts} for field, counts in sorted(field_counts.items())])
    _write_csv(output / "issue_summary.csv", [{"severity": sev, "code": code, "count": count}
                                                 for (sev, code), count in sorted(issue_counts.items())])
    _write_csv(output / "suspicious_rows.csv", suspicious)
    _write_csv(output / "high_confidence_errors.csv", high_errors)
    _write_csv(output / "high_confidence_pass.csv", high_pass)
    (output / "dictionary_candidates.json").write_text(json.dumps(candidates, ensure_ascii=False, indent=2), encoding="utf-8")
    (output / "dictionary_conflicts.json").write_text(json.dumps(conflicts, ensure_ascii=False, indent=2), encoding="utf-8")
    (output / "dictionary_statistics.json").write_text(json.dumps({"candidates": len(candidates), "conflicts": len(conflicts),
        "promoted": 0, "reason": "LOCAL_ZH_NOT_GOLD"}, ensure_ascii=False, indent=2), encoding="utf-8")
    (output / "regression_candidates.json").write_text(json.dumps(high_errors[:300], ensure_ascii=False, indent=2), encoding="utf-8")
    return summary


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--es", required=True, type=Path)
    parser.add_argument("--zh", required=True, type=Path)
    parser.add_argument("--out", type=Path, default=Path("runtime/local_translation_benchmark"))
    args = parser.parse_args(argv)
    if not args.es.is_file() or not args.zh.is_file():
        parser.error("INPUT_NOT_MATERIALIZED")
    print(json.dumps(benchmark(args.es, args.zh, args.out), ensure_ascii=False, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
